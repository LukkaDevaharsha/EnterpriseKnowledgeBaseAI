"""Format-aware extraction for the shared document indexing pipeline."""

import csv
import io
import os
import zipfile

import pymupdf

MAX_ARCHIVE_MEMBERS = 5000
MAX_EXPANDED_SIZE = 100 * 1024 * 1024
MAX_ROWS = 20000


def _read(path):
    if isinstance(path, (bytes, bytearray, memoryview)):
        return bytes(path)
    with open(path, "rb") as file:
        return file.read()


def _check_office_archive(content):
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ARCHIVE_MEMBERS or sum(x.file_size for x in entries) > MAX_EXPANDED_SIZE:
                raise ValueError("The Office document archive exceeds safe limits.")
            if any(x.file_size > 0 and x.file_size / max(x.compress_size, 1) > 1000 for x in entries):
                raise ValueError("The Office document archive exceeds safe compression limits.")
            bad = archive.testzip()
            if bad:
                raise ValueError("The Office document archive is corrupted.")
            return set(archive.namelist())
    except (zipfile.BadZipFile, OSError, RuntimeError) as exc:
        raise ValueError("The Office document is corrupted or invalid.") from exc


def _decode_text(content):
    if b"\x00" in content:
        raise ValueError("The text document contains invalid binary data.")
    encodings = ("utf-16", "utf-8-sig", "cp1252", "latin-1") if content.startswith((b"\xff\xfe", b"\xfe\xff")) else ("utf-8-sig", "cp1252", "latin-1")
    for encoding in encodings:
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("The text document encoding is not supported.")


def _decode_markdown(content):
    if b"\x00" in content:
        raise ValueError("The Markdown document contains invalid binary data.")
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("The Markdown document must use UTF-8 encoding.") from exc


def extract_document_text(file_path, file_type):
    return "\n".join(text for _, text in extract_document_pages(file_path, file_type)).strip()


def extract_document_pages(file_path, file_type):
    """Return ordered passage units; integer references are reserved for PDF pages/slides."""
    kind = (file_type or "").upper()
    try:
        content = _read(file_path)
        if not content:
            raise ValueError("The document is empty.")
        if kind == "PDF":
            if not content.startswith(b"%PDF-"):
                raise ValueError("The uploaded file is not a valid PDF.")
            with pymupdf.open(stream=content, filetype="pdf") as pdf:
                if pdf.is_encrypted:
                    raise ValueError("Password-protected PDF files are not supported.")
                pages = [(i + 1, page.get_text("text").strip()) for i, page in enumerate(pdf)]
        elif kind == "MD":
            pages = [(None, _decode_markdown(content).strip())]
        elif kind == "TXT":
            text = _decode_text(content).strip()
            pages = [(None, text)]
        elif kind in {"DOCX", "PPTX", "XLSX"}:
            names = _check_office_archive(content)
            if kind == "DOCX":
                if "word/document.xml" not in names:
                    raise ValueError("The DOCX document is invalid.")
                from docx import Document as WordDocument
                doc = WordDocument(io.BytesIO(content))
                units = []
                section = ""
                for paragraph in doc.paragraphs:
                    value = paragraph.text.strip()
                    if not value:
                        continue
                    if paragraph.style and paragraph.style.name.lower().startswith("heading"):
                        section = value
                        units.append((None, f"Heading: {value}"))
                    else:
                        units.append((None, f"{section + ' — ' if section else ''}{value}"))
                for table_index, table in enumerate(doc.tables, 1):
                    rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
                    if rows:
                        units.append((None, f"Table {table_index}:\n" + "\n".join(rows)))
                pages = units
            elif kind == "PPTX":
                if "ppt/presentation.xml" not in names:
                    raise ValueError("The PPTX presentation is invalid.")
                from pptx import Presentation
                presentation = Presentation(io.BytesIO(content))
                pages = []
                for number, slide in enumerate(presentation.slides, 1):
                    values = []
                    for shape in slide.shapes:
                        if getattr(shape, "has_text_frame", False) and shape.text.strip():
                            values.append(shape.text.strip())
                        if getattr(shape, "has_table", False):
                            values.extend(" | ".join(cell.text.strip() for cell in row.cells)
                                          for row in shape.table.rows)
                    pages.append((number, f"Slide {number}:\n" + "\n".join(values).strip()))
            else:
                if "xl/workbook.xml" not in names:
                    raise ValueError("The XLSX workbook is invalid.")
                from openpyxl import load_workbook
                workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
                pages = []
                for sheet in workbook.worksheets:
                    rows = []
                    for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                        if row_number > MAX_ROWS:
                            break
                        values = ["" if value is None else str(value) for value in row]
                        if any(values):
                            rows.append(f"Row {row_number}: " + " | ".join(values))
                    pages.append((None, f"Worksheet: {sheet.title}\n" + "\n".join(rows)))
                workbook.close()
        elif kind == "CSV":
            text = _decode_text(content)
            sample = text[:8192]
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
            except csv.Error:
                dialect = csv.excel
            rows = []
            for index, row in enumerate(csv.reader(io.StringIO(text), dialect), 1):
                if index > MAX_ROWS:
                    break
                rows.append(f"Row {index}: " + " | ".join(row))
            pages = [(None, "\n".join(rows))]
        else:
            raise ValueError("Unsupported document type.")
    except ValueError:
        raise
    except (OSError, UnicodeError, RuntimeError, zipfile.BadZipFile, pymupdf.FileDataError,
            pymupdf.EmptyFileError, KeyError, csv.Error) as exc:
        raise ValueError(f"The {kind} document is corrupted or could not be read.") from exc
    if not any(text and text.strip() for _, text in pages):
        raise ValueError("The document contains no extractable text.")
    return pages


def extract_pdf_outline(file_path):
    try:
        with pymupdf.open(stream=_read(file_path), filetype="pdf") as pdf:
            return [{"level": level, "title": title, "page_number": number}
                    for level, title, number, *_ in pdf.get_toc()
                    if isinstance(level, int) and isinstance(title, str)
                    and isinstance(number, int) and number > 0]
    except (pymupdf.FileDataError, pymupdf.EmptyFileError, OSError) as exc:
        raise ValueError("The PDF outline could not be read.") from exc
